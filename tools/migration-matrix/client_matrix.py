#!/usr/bin/env python3
"""Client content-matrix ingest — Phase 1's second half.

Crystal's per-client matrix (one row per page, lifecycle columns across
disposition -> audit -> writing -> migration x3 rounds -> QA -> client
review) feeds the same master database as Laura's tracker. Start with one
matrix, prove it, add more: this module globs for *content-matrix*.xlsx
files in the parent folder and ingests every one it finds."""

import glob
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))

# 2-programs sheet layout (1-indexed columns; headers on rows 2-3, data 6+)
COL = {
    "url": 1, "assigned_to": 2, "page_id": 3, "title": 4, "new_url": 5,
    "disp_delete": 6, "disp_reuse": 7, "disp_write_new": 8, "disp_revise": 9,
    "disp_optimize": 10, "disp_as_is": 11, "disp_import": 12,
    "writer": 22, "draft_delivered": 27, "reviewer": 28, "r1_review": 29,
    "r1_updates": 30,
    "mig_assigned_to": 37, "mig_r1_by": 41, "mig_r1_done": 42,
    "qa_r1_by": 44, "qa_r1_done": 45, "mig_r2_done": 47, "qa_r2_done": 49,
    "delivered": 54, "client_final": 61, "mig_r3_done": 51,
}
DISPOSITIONS = ["Write New", "Revise", "Optimize", "Reuse", "As Is",
                "Import", "Delete"]
DISP_COLS = ["disp_write_new", "disp_revise", "disp_optimize", "disp_reuse",
             "disp_as_is", "disp_import", "disp_delete"]
# funnel: a page has "reached" a stage if that cell is filled
FUNNEL = [
    ("Draft delivered", "draft_delivered"),
    ("Writing review done", "r1_updates"),
    ("Migrated (R1)", "mig_r1_done"),
    ("QA passed (R1)", "qa_r1_done"),
    ("Delivered to client", "delivered"),
    ("Client final review", "client_final"),
]


def _val(ws, r, key):
    v = ws.cell(r, COL[key]).value
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def find_matrices():
    # MM_MATRIX_DIR points the glob at a synced OneDrive/SharePoint folder;
    # default stays the project folder (the parent of this script's dir).
    parent = os.environ.get("MM_MATRIX_DIR") or \
        os.path.abspath(os.path.join(HERE, ".."))
    return sorted(p for p in glob.glob(os.path.join(parent, "*.xlsx"))
                  if "content-matrix" in os.path.basename(p).lower()
                  and not os.path.basename(p).startswith("~$")
                  and "-FIXED" not in os.path.basename(p))  # our own output


def parse_matrix(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    name = os.path.basename(path).replace(".xlsx", "")
    out = {"name": name, "sections": [], "pages": [], "broken_refs": 0}
    # their own DASHBOARD rollup, if present, and whether it has rotted
    for ws in wb.worksheets:
        if ws.title.upper() == "DASHBOARD":
            out["broken_refs"] = sum(
                1 for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and "#REF!" in c.value)
            continue
        if ws.max_column < 60:  # not a page-lifecycle sheet
            continue
        out["sections"].append(ws.title)
        for r in range(6, ws.max_row + 1):
            if not _val(ws, r, "url"):
                continue
            page = {k: _val(ws, r, k) for k in COL}
            disp = next((lab for lab, key in zip(DISPOSITIONS, DISP_COLS)
                         if page[key]), None)
            stage = "Not started"
            for lab, key in FUNNEL:
                if page[key]:
                    stage = lab
            out["pages"].append({
                "section": ws.title, "url": page["url"],
                "page_id": page["page_id"], "title": page["title"],
                "disposition": disp, "writer": page["writer"],
                "reviewer": page["reviewer"],
                "migrator": page["mig_r1_by"], "qa": page["qa_r1_by"],
                "stage": stage,
                "draft_date": (page["draft_delivered"] or "")[:10] or None,
                "delivered_date": (page["delivered"] or "")[:10] or None,
                "rework": bool(page["mig_r2_done"] or page["mig_r3_done"]),
                **{key: bool(page[key]) for _, key in FUNNEL},
            })
    return out


def summarize(m):
    pages = m["pages"]
    total = len(pages)
    disp = {}
    for p in pages:
        disp[p["disposition"] or "(none)"] = disp.get(p["disposition"] or "(none)", 0) + 1
    funnel = [(lab, sum(1 for p in pages if p[key])) for lab, key in FUNNEL]
    return {
        "name": m["name"], "sections": m["sections"], "total": total,
        "dispositions": disp, "funnel": funnel,
        "no_disposition": disp.get("(none)", 0),
        "broken_refs": m["broken_refs"],
    }


def ingest(db):
    """Load every client matrix into the master DB. Returns summaries."""
    db.executescript("""
        DROP TABLE IF EXISTS client_pages;
        CREATE TABLE client_pages (matrix TEXT, section TEXT, url TEXT,
            page_id TEXT, title TEXT, disposition TEXT, writer TEXT,
            reviewer TEXT, migrator TEXT, qa TEXT, stage TEXT,
            draft_delivered INT, writing_done INT, migrated INT,
            qa_done INT, delivered INT, client_final INT,
            draft_date TEXT, delivered_date TEXT, rework INT);
    """)
    summaries = []
    for path in find_matrices():
        m = parse_matrix(path)
        for p in m["pages"]:
            db.execute(
                "INSERT INTO client_pages VALUES "
                "(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (m["name"], p["section"], p["url"], p["page_id"], p["title"],
                 p["disposition"], p["writer"], p["reviewer"], p["migrator"],
                 p["qa"], p["stage"], p["draft_delivered"], p["r1_updates"],
                 p["mig_r1_done"], p["qa_r1_done"], p["delivered"],
                 p["client_final"], p["draft_date"], p["delivered_date"],
                 p["rework"]))
        summaries.append(summarize(m))
    db.commit()
    return summaries


def link_to_project(db, summary):
    """Best-effort link: the matrix's dominant client word (from its
    Assigned Reviewer column) vs tracker project names — e.g. reviewer
    'Washtenaw' -> project 'Washtenaw Community College'."""
    row = db.execute("""
        SELECT reviewer, COUNT(*) FROM client_pages
        WHERE matrix=? AND reviewer IS NOT NULL AND reviewer != 'Stamats'
        GROUP BY reviewer ORDER BY COUNT(*) DESC LIMIT 1
    """, (summary["name"],)).fetchone()
    hint = row[0] if row else None
    projects = list(db.execute(
        "SELECT name, total_pages FROM projects WHERE name IS NOT NULL"))
    if hint:
        for pname, total in projects:
            if hint.lower().split()[0] in (pname or "").lower():
                return {"project": pname, "hint": hint, "total_pages": total}
    # fallback: match a filename token against project names
    # (e.g. content-matrix-morehead -> "Morehead Alumni Foundation")
    for token in summary["name"].lower().replace("_", "-").split("-"):
        if len(token) < 5 or token in ("content", "matrix", "client", "sample"):
            continue
        for pname, total in projects:
            if token in (pname or "").lower():
                return {"project": pname, "hint": f"filename '{token}'",
                        "total_pages": total}
    if hint:
        return {"project": None, "hint": hint, "total_pages": None}
    return None
